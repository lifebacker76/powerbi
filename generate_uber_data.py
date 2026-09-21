"""
Generate synthetic ride-sharing dataset for the Uber Analytics case study.
Produces an Excel workbook with 6 tables matching the case study spec exactly.

Data quality issues are INTENTIONALLY included in the raw data:
  - ~8% of completed trips have missing tip values
  - ~6% of PaymentMethod values have inconsistent casing
  - Cancelled trips carry 0 for fare/distance/duration
"""

import random
import datetime
import uuid
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter

random.seed(2025)

# ── Constants ───────────────────────────────────────────────────────────────

CITIES_ZONES = {
    "Mumbai": ["Andheri", "Bandra", "Churchgate", "Malad", "Thane", "Powai", "Dadar"],
    "Bangalore": ["Koramangala", "Whitefield", "MG Road", "Electronic City", "Indiranagar", "HSR Layout", "Marathahalli"],
    "Delhi": ["Connaught Place", "Dwarka", "Saket", "Karol Bagh", "Lajpat Nagar", "Rohini", "Hauz Khas"],
    "Hyderabad": ["HITEC City", "Banjara Hills", "Gachibowli", "Secunderabad", "Ameerpet", "Jubilee Hills"],
    "Chennai": ["T Nagar", "Anna Nagar", "Adyar", "OMR", "Nungambakkam", "Velachery", "Guindy"],
    "Kolkata": ["Park Street", "Salt Lake", "Howrah", "New Town", "Ballygunge", "Alipore"],
    "Pune": ["Hinjewadi", "Koregaon Park", "Kothrud", "Viman Nagar", "Hadapsar", "Baner", "Aundh"],
    "Ahmedabad": ["SG Highway", "Navrangpura", "Satellite", "Vastrapur", "Prahlad Nagar", "Bodakdev"],
    "Jaipur": ["MI Road", "Malviya Nagar", "Mansarovar", "C-Scheme", "Vaishali Nagar", "Tonk Road"],
    "Lucknow": ["Hazratganj", "Gomti Nagar", "Aliganj", "Indira Nagar", "Mahanagar", "Alambagh"],
    "Chandigarh": ["Sector 17", "Sector 22", "Sector 35", "IT Park", "Industrial Area"],
    "Kochi": ["MG Road Kochi", "Edappally", "Kakkanad", "Fort Kochi", "Vyttila", "Aluva"],
}

VEHICLE_TYPES = ["Auto", "Mini", "Sedan", "SUV", "Premium"]
VEHICLE_WEIGHTS = [0.15, 0.30, 0.30, 0.15, 0.10]

LOYALTY_TIERS = ["Bronze", "Silver", "Gold", "Platinum"]
LOYALTY_WEIGHTS = [0.45, 0.30, 0.18, 0.07]

PAYMENT_METHODS_CLEAN = ["UPI", "Cash", "Card", "Wallet"]
PAYMENT_WEIGHTS = [0.35, 0.25, 0.25, 0.15]
PAYMENT_DIRTY_VARIANTS = {
    "UPI": ["upi", "Upi", "UPI"],
    "Cash": ["cash", "CASH", "Cash"],
    "Card": ["card", "CARD", "Card"],
    "Wallet": ["wallet", "WALLET", "Wallet"],
}

CANCEL_REASONS = [
    "Long wait time",
    "High surge pricing",
    "Found alternative",
    "Driver cancelled",
    "Changed plans",
    "Other",
]
CANCEL_REASON_WEIGHTS = [0.30, 0.25, 0.15, 0.15, 0.10, 0.05]

PROMOS = [
    {"PromoKey": 1, "PromoCode": "FIRST50",   "CampaignName": "New User Welcome",      "DiscountPct": 50},
    {"PromoKey": 2, "PromoCode": "WEEKEND20",  "CampaignName": "Weekend Rides",          "DiscountPct": 20},
    {"PromoKey": 3, "PromoCode": "LOYALTY15",  "CampaignName": "Loyalty Rewards",        "DiscountPct": 15},
    {"PromoKey": 4, "PromoCode": "MONSOON10",  "CampaignName": "Monsoon Special",        "DiscountPct": 10},
    {"PromoKey": 5, "PromoCode": "CORP25",     "CampaignName": "Corporate Program",      "DiscountPct": 25},
    {"PromoKey": 6, "PromoCode": "FESTIVE30",  "CampaignName": "Festive Season Bonanza", "DiscountPct": 30},
]

FIRST_NAMES = [
    "Aarav","Vivaan","Aditya","Vihaan","Arjun","Sai","Reyansh","Ayaan","Krishna","Ishaan",
    "Ananya","Aadhya","Myra","Sara","Aanya","Priya","Neha","Riya","Isha","Kavya",
    "Rohan","Rahul","Amit","Vikram","Raj","Suresh","Mahesh","Karan","Nikhil","Akash",
    "Pooja","Anjali","Deepika","Shreya","Tanvi","Meera","Divya","Nisha","Sneha","Swati",
    "Manish","Rajesh","Sunil","Arun","Vijay","Sanjay","Ravi","Gaurav","Pankaj","Deepak",
]
LAST_NAMES = [
    "Sharma","Verma","Gupta","Singh","Kumar","Patel","Reddy","Nair","Iyer","Joshi",
    "Mehta","Shah","Das","Roy","Mukherjee","Chatterjee","Banerjee","Sen","Pillai","Menon",
    "Rao","Naidu","Hegde","Kulkarni","Deshmukh","Patil","Jain","Agarwal","Saxena","Mishra",
]

# ── City demand weights (Mumbai & Bangalore are highest-demand) ─────────
CITY_TRIP_WEIGHTS = {
    "Mumbai": 0.18, "Bangalore": 0.17, "Delhi": 0.14, "Hyderabad": 0.10,
    "Chennai": 0.09, "Kolkata": 0.07, "Pune": 0.07, "Ahmedabad": 0.05,
    "Jaipur": 0.04, "Lucknow": 0.04, "Chandigarh": 0.03, "Kochi": 0.02,
}

# ── Helpers ──────────────────────────────────────────────────────────────

def weighted_choice(options, weights):
    return random.choices(options, weights=weights, k=1)[0]


def generate_surge(hour, city):
    """Generate surge multiplier based on time and city."""
    is_peak = 17 <= hour <= 21
    is_morning_peak = 7 <= hour <= 10
    is_high_demand_city = city in ("Mumbai", "Bangalore")

    if is_peak and is_high_demand_city:
        r = random.random()
        if r < 0.25:
            return round(random.uniform(2.0, 3.0), 1)
        elif r < 0.55:
            return round(random.uniform(1.5, 2.0), 1)
        elif r < 0.80:
            return round(random.uniform(1.1, 1.5), 1)
        else:
            return 1.0
    elif is_peak:
        r = random.random()
        if r < 0.10:
            return round(random.uniform(2.0, 2.8), 1)
        elif r < 0.35:
            return round(random.uniform(1.5, 2.0), 1)
        elif r < 0.60:
            return round(random.uniform(1.1, 1.5), 1)
        else:
            return 1.0
    elif is_morning_peak:
        r = random.random()
        if r < 0.05:
            return round(random.uniform(2.0, 2.5), 1)
        elif r < 0.20:
            return round(random.uniform(1.5, 2.0), 1)
        elif r < 0.45:
            return round(random.uniform(1.1, 1.5), 1)
        else:
            return 1.0
    else:
        r = random.random()
        if r < 0.02:
            return round(random.uniform(2.0, 2.3), 1)
        elif r < 0.08:
            return round(random.uniform(1.5, 2.0), 1)
        elif r < 0.22:
            return round(random.uniform(1.1, 1.5), 1)
        else:
            return 1.0


def should_cancel(surge, hour, city):
    """Determine if a trip is cancelled based on surge, time, city."""
    base = 0.04
    if surge >= 2.0:
        base = 0.24
    elif surge >= 1.5:
        base = 0.11
    elif surge >= 1.1:
        base = 0.055

    if 17 <= hour <= 21 and city in ("Mumbai", "Bangalore") and surge >= 2.0:
        base = min(base + 0.08, 0.45)
    elif 17 <= hour <= 21 and surge >= 1.5:
        base = min(base + 0.03, 0.35)

    return random.random() < base


def dirty_payment(method):
    """~6% chance of returning an inconsistently-cased payment method."""
    if random.random() < 0.06:
        variants = PAYMENT_DIRTY_VARIANTS[method]
        return random.choice([v for v in variants if v != method] or variants)
    return method


# ── Build Dim_Location ──────────────────────────────────────────────────

print("Building Dim_Location...")
locations = []
loc_key = 1
city_locations = {}
for city, zones in CITIES_ZONES.items():
    city_locations[city] = []
    for zone in zones:
        locations.append({
            "LocationKey": loc_key,
            "City": city,
            "Zone": zone,
        })
        city_locations[city].append(loc_key)
        loc_key += 1

print(f"  → {len(locations)} locations across {len(CITIES_ZONES)} cities")

# ── Build Dim_Date ──────────────────────────────────────────────────────

print("Building Dim_Date...")
start_date = datetime.date(2025, 1, 1)
end_date = datetime.date(2026, 12, 31)
dates = []
d = start_date
date_key_map = {}
while d <= end_date:
    dk = int(d.strftime("%Y%m%d"))
    dates.append({
        "DateKey": dk,
        "Date": d,
        "Year": d.year,
        "Quarter": f"Q{(d.month - 1) // 3 + 1}",
        "MonthName": d.strftime("%B"),
        "MonthNum": d.month,
        "Day": d.day,
        "DayOfWeek": d.strftime("%A"),
        "WeekNum": d.isocalendar()[1],
        "IsWeekend": d.weekday() >= 5,
    })
    date_key_map[d] = dk
    d += datetime.timedelta(days=1)

print(f"  → {len(dates)} days")

# ── Build Dim_Driver ────────────────────────────────────────────────────

print("Building Dim_Driver (2,500)...")
drivers = []
cities_list = list(CITIES_ZONES.keys())
city_weights_list = [CITY_TRIP_WEIGHTS[c] for c in cities_list]

for i in range(1, 2501):
    home_city = weighted_choice(cities_list, city_weights_list)
    join_date = start_date - datetime.timedelta(days=random.randint(30, 1500))
    rating = round(random.triangular(3.0, 5.0, 4.5), 1)
    rating = min(5.0, max(1.0, rating))
    drivers.append({
        "DriverKey": i,
        "DriverName": f"{random.choice(FIRST_NAMES)} {random.choice(LAST_NAMES)}",
        "VehicleType": weighted_choice(VEHICLE_TYPES, VEHICLE_WEIGHTS),
        "Rating": rating,
        "JoinDate": join_date,
        "HomeCity": home_city,
    })

print(f"  → {len(drivers)} drivers")

driver_by_city = {}
for drv in drivers:
    driver_by_city.setdefault(drv["HomeCity"], []).append(drv["DriverKey"])

# ── Build Dim_Rider ─────────────────────────────────────────────────────

print("Building Dim_Rider (6,000)...")
riders = []
for i in range(1, 6001):
    signup = start_date - datetime.timedelta(days=random.randint(0, 1200))
    riders.append({
        "RiderKey": i,
        "RiderName": f"{random.choice(FIRST_NAMES)} {random.choice(LAST_NAMES)}",
        "SignupDate": signup,
        "LoyaltyTier": weighted_choice(LOYALTY_TIERS, LOYALTY_WEIGHTS),
    })

print(f"  → {len(riders)} riders")

# ── Build Dim_Promo ─────────────────────────────────────────────────────

print("Building Dim_Promo...")
promos = PROMOS[:]
print(f"  → {len(promos)} promos")

# ── Build Fact_Trips (120,000) ──────────────────────────────────────────

print("Building Fact_Trips (120,000)... this takes a moment.")

all_dates = [start_date + datetime.timedelta(days=i) for i in range((end_date - start_date).days + 1)]
trips = []

FARE_BASE = {"Auto": 30, "Mini": 50, "Sedan": 80, "SUV": 120, "Premium": 180}
FARE_PER_KM = {"Auto": 8, "Mini": 10, "Sedan": 14, "SUV": 18, "Premium": 25}

completed_count = 0
cancelled_count = 0

for trip_id in range(1, 120001):
    city = weighted_choice(cities_list, city_weights_list)
    loc_key_val = random.choice(city_locations[city])

    trip_date = random.choice(all_dates)
    date_key = date_key_map[trip_date]

    hour = weighted_choice(
        list(range(24)),
        [1,1,1,1,1,2, 4,6,7,6,4,3, 3,3,3,4,5,7, 8,8,6,4,2,1]
    )

    drv_pool = driver_by_city.get(city, list(range(1, 101)))
    driver_key = random.choice(drv_pool)
    rider_key = random.choice(range(1, 6001))

    drv = drivers[driver_key - 1]
    vehicle = drv["VehicleType"]

    surge = generate_surge(hour, city)
    is_cancelled = should_cancel(surge, hour, city)

    if random.random() < 0.12:
        promo_key = random.choice([p["PromoKey"] for p in promos])
    else:
        promo_key = 0

    if is_cancelled:
        cancelled_count += 1
        fare = 0.0
        distance = 0.0
        duration = 0
        wait_time = round(random.uniform(3, 25), 1)
        tip = 0.0
        status = "Cancelled"
        cancel_reason = weighted_choice(CANCEL_REASONS, CANCEL_REASON_WEIGHTS)
        payment = weighted_choice(PAYMENT_METHODS_CLEAN, PAYMENT_WEIGHTS)
        payment = dirty_payment(payment)
    else:
        completed_count += 1
        distance = round(random.triangular(1.5, 45, 8), 1)
        duration = max(5, int(distance * random.uniform(2.5, 5.0) + random.uniform(2, 15)))
        base = FARE_BASE[vehicle]
        per_km = FARE_PER_KM[vehicle]
        fare = round((base + distance * per_km) * surge, 2)

        if promo_key > 0:
            promo = next(p for p in promos if p["PromoKey"] == promo_key)
            discount = promo["DiscountPct"] / 100.0
            fare = round(fare * (1 - discount), 2)
            fare = max(fare, 20.0)

        wait_time = round(random.triangular(1, 15, 4), 1)

        if random.random() < 0.08:
            tip = None
        elif random.random() < 0.60:
            tip = round(random.uniform(5, max(fare * 0.15, 10)), 0)
        else:
            tip = 0.0

        status = "Completed"
        cancel_reason = None
        payment = weighted_choice(PAYMENT_METHODS_CLEAN, PAYMENT_WEIGHTS)
        payment = dirty_payment(payment)

    trips.append({
        "TripKey": trip_id,
        "DateKey": date_key,
        "DriverKey": driver_key,
        "RiderKey": rider_key,
        "LocationKey": loc_key_val,
        "PromoKey": promo_key,
        "HourOfDay": hour,
        "Fare": fare,
        "DistanceKM": distance,
        "DurationMin": duration,
        "WaitTimeMin": wait_time,
        "Tip": tip,
        "SurgeMultiplier": surge,
        "Status": status,
        "CancellationReason": cancel_reason,
        "PaymentMethod": payment,
    })

    if trip_id % 30000 == 0:
        print(f"  ... {trip_id:,} trips generated")

total = completed_count + cancelled_count
cancel_rate = cancelled_count / total * 100
high_surge = [t for t in trips if t["SurgeMultiplier"] >= 2.0]
high_surge_cancel = sum(1 for t in high_surge if t["Status"] == "Cancelled")
high_surge_rate = high_surge_cancel / len(high_surge) * 100 if high_surge else 0

print(f"  → {total:,} trips ({completed_count:,} completed, {cancelled_count:,} cancelled)")
print(f"  → Overall cancellation rate: {cancel_rate:.1f}%")
print(f"  → High-surge (≥2.0) cancellation rate: {high_surge_rate:.1f}%")

# ── Write to Excel ──────────────────────────────────────────────────────

print("\nWriting Excel workbook...")

HEADER_FONT = Font(bold=True, color="FFFFFF", size=10, name="Segoe UI")
HEADER_FILL = PatternFill(start_color="1B1B2F", end_color="1B1B2F", fill_type="solid")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center")
THIN_BORDER = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)
ALT_FILL = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")


def style_sheet(ws, num_cols):
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, max_col=num_cols, max_row=ws.max_row), start=2):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center")
            cell.font = Font(name="Segoe UI", size=9)
            if row_idx % 2 == 0:
                cell.fill = ALT_FILL
    ws.freeze_panes = "A2"
    for col in range(1, num_cols + 1):
        max_len = max(len(str(c.value or "")) for c in ws[get_column_letter(col)])
        ws.column_dimensions[get_column_letter(col)].width = min(max_len + 3, 30)


def write_table(wb, name, headers, rows):
    ws = wb.create_sheet(title=name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h) for h in headers])
    style_sheet(ws, len(headers))
    return ws


wb = openpyxl.Workbook()

trip_headers = list(trips[0].keys())
write_table(wb, "Fact_Trips", trip_headers, trips)

driver_headers = list(drivers[0].keys())
write_table(wb, "Dim_Driver", driver_headers, drivers)

rider_headers = list(riders[0].keys())
write_table(wb, "Dim_Rider", rider_headers, riders)

loc_headers = list(locations[0].keys())
write_table(wb, "Dim_Location", loc_headers, locations)

date_headers = list(dates[0].keys())
write_table(wb, "Dim_Date", date_headers, dates)

promo_headers = list(promos[0].keys())
write_table(wb, "Dim_Promo", promo_headers, promos)

if "Sheet" in wb.sheetnames:
    del wb["Sheet"]

OUTPUT = "/Users/bsinga1/Desktop/Powerbi/RideShareAnalytics_Data.xlsx"
wb.save(OUTPUT)

import os
size_mb = os.path.getsize(OUTPUT) / (1024 * 1024)
print(f"\n{'='*60}")
print(f"✅  Workbook saved → {OUTPUT}")
print(f"    Size: {size_mb:.1f} MB")
print(f"    Sheets: {wb.sheetnames}")
print(f"{'='*60}")
