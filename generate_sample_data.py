"""
Generate a sample Excel workbook with dummy data for Power BI.
Tables: Sales, Products, Customers, Stores, Calendar
Relationships are modeled via shared key columns so Power BI can auto-detect them.
"""

import random
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

random.seed(42)

# ── Helper ──────────────────────────────────────────────────────────────────
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center")
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


def style_sheet(ws, num_cols):
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER
    for row in ws.iter_rows(min_row=2, max_col=num_cols, max_row=ws.max_row):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center")
    for col in range(1, num_cols + 1):
        max_len = max(len(str(c.value or "")) for c in ws[get_column_letter(col)])
        ws.column_dimensions[get_column_letter(col)].width = max_len + 4


# ── Reference data ──────────────────────────────────────────────────────────
CATEGORIES = ["Electronics", "Clothing", "Home & Kitchen", "Sports", "Books"]
SUB_CATEGORIES = {
    "Electronics": ["Laptops", "Phones", "Tablets", "Headphones", "Cameras"],
    "Clothing": ["Shirts", "Pants", "Jackets", "Shoes", "Accessories"],
    "Home & Kitchen": ["Furniture", "Appliances", "Cookware", "Decor", "Lighting"],
    "Sports": ["Fitness", "Outdoor", "Team Sports", "Water Sports", "Cycling"],
    "Books": ["Fiction", "Non-Fiction", "Science", "History", "Technology"],
}
CITIES = [
    ("New York", "NY", "East"),
    ("Los Angeles", "CA", "West"),
    ("Chicago", "IL", "Central"),
    ("Houston", "TX", "South"),
    ("Phoenix", "AZ", "West"),
    ("Philadelphia", "PA", "East"),
    ("San Antonio", "TX", "South"),
    ("Dallas", "TX", "South"),
    ("San Jose", "CA", "West"),
    ("Seattle", "WA", "West"),
]
FIRST_NAMES = ["James", "Mary", "Robert", "Patricia", "John", "Jennifer",
               "Michael", "Linda", "David", "Elizabeth", "William", "Barbara",
               "Richard", "Susan", "Joseph", "Jessica", "Thomas", "Sarah",
               "Charles", "Karen"]
LAST_NAMES = ["Smith", "Johnson", "Williams", "Brown", "Jones", "Garcia",
              "Miller", "Davis", "Rodriguez", "Martinez", "Hernandez", "Lopez",
              "Gonzalez", "Wilson", "Anderson", "Thomas", "Taylor", "Moore",
              "Jackson", "Martin"]
SEGMENTS = ["Consumer", "Corporate", "Small Business"]
CHANNELS = ["Online", "In-Store", "Phone"]

# ── Build Products ──────────────────────────────────────────────────────────
products = []
pid = 1000
for cat in CATEGORIES:
    for sub in SUB_CATEGORIES[cat]:
        for variant in range(1, 4):
            products.append({
                "ProductID": f"P-{pid}",
                "ProductName": f"{sub} Pro {variant}",
                "Category": cat,
                "SubCategory": sub,
                "UnitCost": round(random.uniform(5, 200), 2),
                "UnitPrice": 0,
            })
            products[-1]["UnitPrice"] = round(
                products[-1]["UnitCost"] * random.uniform(1.3, 2.5), 2
            )
            pid += 1

# ── Build Stores ────────────────────────────────────────────────────────────
stores = []
for i, (city, state, region) in enumerate(CITIES, start=1):
    stores.append({
        "StoreID": f"S-{100 + i}",
        "StoreName": f"{city} Store",
        "City": city,
        "State": state,
        "Region": region,
    })

# ── Build Customers ─────────────────────────────────────────────────────────
customers = []
for i in range(1, 51):
    city, state, region = random.choice(CITIES)
    customers.append({
        "CustomerID": f"C-{2000 + i}",
        "FirstName": random.choice(FIRST_NAMES),
        "LastName": random.choice(LAST_NAMES),
        "Email": "",
        "Segment": random.choice(SEGMENTS),
        "City": city,
        "State": state,
    })
    c = customers[-1]
    c["Email"] = f"{c['FirstName'].lower()}.{c['LastName'].lower()}{i}@example.com"

# ── Build Calendar (2 years) ───────────────────────────────────────────────
start_date = datetime.date(2024, 1, 1)
end_date = datetime.date(2025, 12, 31)
calendar_rows = []
d = start_date
while d <= end_date:
    calendar_rows.append({
        "Date": d,
        "Year": d.year,
        "Quarter": f"Q{(d.month - 1) // 3 + 1}",
        "Month": d.strftime("%B"),
        "MonthNum": d.month,
        "Day": d.day,
        "DayOfWeek": d.strftime("%A"),
        "IsWeekend": d.weekday() >= 5,
    })
    d += datetime.timedelta(days=1)

# ── Build Sales Transactions ────────────────────────────────────────────────
sales = []
order_id = 5000
for _ in range(500):
    order_id += 1
    order_date = start_date + datetime.timedelta(
        days=random.randint(0, (end_date - start_date).days)
    )
    cust = random.choice(customers)
    store = random.choice(stores)
    prod = random.choice(products)
    qty = random.randint(1, 10)
    discount = random.choice([0, 0, 0, 0.05, 0.10, 0.15, 0.20])
    unit_price = prod["UnitPrice"]
    total = round(qty * unit_price * (1 - discount), 2)
    cost = round(qty * prod["UnitCost"], 2)
    profit = round(total - cost, 2)

    sales.append({
        "OrderID": f"ORD-{order_id}",
        "OrderDate": order_date,
        "CustomerID": cust["CustomerID"],
        "ProductID": prod["ProductID"],
        "StoreID": store["StoreID"],
        "Channel": random.choice(CHANNELS),
        "Quantity": qty,
        "UnitPrice": unit_price,
        "Discount": discount,
        "TotalAmount": total,
        "Cost": cost,
        "Profit": profit,
    })

# ── Write workbook ──────────────────────────────────────────────────────────
wb = openpyxl.Workbook()

def write_table(wb, name, headers, rows):
    ws = wb.create_sheet(title=name)
    ws.append(headers)
    for row in rows:
        ws.append([row[h] for h in headers])
    style_sheet(ws, len(headers))
    return ws

write_table(wb, "Sales", list(sales[0].keys()), sales)
write_table(wb, "Products", list(products[0].keys()), products)
write_table(wb, "Customers", list(customers[0].keys()), customers)
write_table(wb, "Stores", list(stores[0].keys()), stores)
write_table(wb, "Calendar", list(calendar_rows[0].keys()), calendar_rows)

# Remove default empty sheet
if "Sheet" in wb.sheetnames:
    del wb["Sheet"]

OUTPUT = "/Users/bsinga1/Desktop/Powerbi/PowerBI_Sample_Data.xlsx"
wb.save(OUTPUT)
print(f"✅  Workbook saved → {OUTPUT}")
print(f"    Sheets : {wb.sheetnames}")
print(f"    Sales  : {len(sales):,} rows")
print(f"    Products: {len(products):,} rows")
print(f"    Customers: {len(customers):,} rows")
print(f"    Stores : {len(stores):,} rows")
print(f"    Calendar: {len(calendar_rows):,} rows")
